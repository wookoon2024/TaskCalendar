from __future__ import annotations

import json
import logging
import os
import re
import subprocess
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from PySide6.QtCore import QEvent, QObject, QPoint, QSize, QStringListModel, Qt, QTimer, QUrl
from PySide6.QtGui import QBrush, QColor, QCursor, QDesktopServices, QFont, QIcon
from PySide6.QtWidgets import (
    QApplication,
    QCheckBox,
    QComboBox,
    QCompleter,
    QDateEdit,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFormLayout,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLayout,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMenu,
    QMessageBox,
    QPushButton,
    QSizePolicy,
    QSpinBox,
    QStyle,
    QStyledItemDelegate,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from taskcalendar.models import CalendarEntry, EntryType
from taskcalendar.paths import asset_path
from taskcalendar.storage import EncryptedRepository
from taskcalendar.qt_styles import resolve_palette, _shade
from taskcalendar.themes import THEMES

logger = logging.getLogger(__name__)


def get_dialog_palette(parent=None, repository=None, main_window=None) -> dict[str, str]:
    """부모 위젯, 메인 윈도우, 저장소 설정에서 현재 활성 스킨 팔레트를 안전하게 추출"""
    if parent:
        p = getattr(parent, "palette", None)
        if isinstance(p, dict) and p.get("bg"):
            return p
    if main_window:
        p = getattr(main_window, "palette", None)
        if isinstance(p, dict) and p.get("bg"):
            return p
    if repository:
        t = repository.get_setting("theme", "light")
        if t in THEMES:
            return THEMES[t]
    return THEMES["light"]


class SafeDateEdit(QDateEdit):
    """
    마우스 휠 스크롤, 트랙패드 제스처, 더블클릭, 스핀박스 스텝으로 인해 날짜(년도)가
    의도치 않게 급변(예: 2026 -> 2027 -> 2028)하는 현상을 완벽 차단하는 QDateEdit
    """

    def __init__(self, *args, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self.setCalendarPopup(True)
        le = self.lineEdit()
        if le:
            le.installEventFilter(self)
        cal = self.calendarWidget()
        if cal:
            cal.installEventFilter(self)
            year_edit = cal.findChild(QSpinBox, "qt_calendar_yearedit")
            if year_edit:
                year_edit.installEventFilter(self)

    def stepBy(self, steps: int) -> None:
        # 스핀박스 증감에 의한 연도/날짜 변동 차단 (2026->2027 방지)
        pass

    def wheelEvent(self, event) -> None:
        # 마우스 휠에 의한 날짜/연도 증감 차단
        event.ignore()

    def mouseDoubleClickEvent(self, event) -> None:
        # 더블클릭 시 텍스트 전체 선택 (연도 증감 차단)
        le = self.lineEdit()
        if le:
            le.selectAll()
        event.accept()

    def eventFilter(self, watched: QObject, event: QEvent) -> bool:
        if event.type() == QEvent.Type.Wheel:
            event.ignore()
            return True
        if event.type() == QEvent.Type.MouseButtonDblClick:
            le = self.lineEdit()
            if le:
                le.selectAll()
            event.accept()
            return True
        return super().eventFilter(watched, event)


HEADER_BUTTON_STYLE = """
    QPushButton {
        background-color: #FFFFFF;
        color: #1F2328;
        border: 1px solid #CBD5E0;
        border-radius: 6px;
        padding: 0 12px;
        font-size: 12px;
    }
    QPushButton:hover {
        background-color: #F6F8FB;
        border-color: #A0AEC0;
    }
"""

PRIMARY_BUTTON_STYLE = """
    QPushButton {
        background-color: #1F7A67;
        color: #FFFFFF;
        border: 1px solid #1F7A67;
        border-radius: 6px;
        padding: 0 12px;
        font-size: 12px;
        font-weight: 600;
    }
    QPushButton:hover {
        background-color: #185F50;
        border-color: #185F50;
    }
    QPushButton:pressed {
        background-color: #144E42;
        border-color: #144E42;
    }
"""

DEFAULT_CATEGORIES = ["일반", "공문", "보고", "기안", "결재", "민원", "예산", "기타"]
DEFAULT_STATUSES = ["등록", "진행중", "완료", "보류", "취소"]


def get_task_categories(repo: EncryptedRepository) -> list[str]:
    cats = list(DEFAULT_CATEGORIES)
    saved = repo.get_setting("custom_task_categories")
    if saved:
        try:
            loaded = json.loads(saved)
            if isinstance(loaded, list) and loaded:
                cats = loaded
        except Exception:
            pass
    # '일반'은 항상 최상단 0번에 고정
    cats = [c for c in cats if c != "일반"]
    return ["일반"] + cats


def save_task_categories(repo: EncryptedRepository, cats: list[str]) -> None:
    # '일반'은 항상 최상단 0번에 고정하여 저장
    clean_cats = [c for c in cats if c != "일반"]
    final_cats = ["일반"] + clean_cats
    repo.set_setting("custom_task_categories", json.dumps(final_cats, ensure_ascii=False))
    repo.save()


def get_task_statuses(repo: EncryptedRepository) -> list[str]:
    statuses = list(DEFAULT_STATUSES)
    saved = repo.get_setting("custom_task_statuses")
    if saved:
        try:
            loaded = json.loads(saved)
            if isinstance(loaded, list) and loaded:
                statuses = loaded
        except Exception:
            pass
    # '등록'은 항상 최상단 0번에 고정
    statuses = [s for s in statuses if s != "등록"]
    return ["등록"] + statuses


def save_task_statuses(repo: EncryptedRepository, statuses: list[str]) -> None:
    # '등록'은 항상 최상단 0번에 고정하여 저장
    clean_statuses = [s for s in statuses if s != "등록"]
    final_statuses = ["등록"] + clean_statuses
    repo.set_setting("custom_task_statuses", json.dumps(final_statuses, ensure_ascii=False))
    repo.save()


STATUS_CONFIG = {
    "등록": {"bg": "#EBF8FF", "text": "#2B6CB0", "border": "#BEE3F8", "icon": "🔵"},
    "진행중": {"bg": "#FFFAF0", "text": "#C05621", "border": "#FEEBC8", "icon": "🟠"},
    "완료": {"bg": "#F0FFF4", "text": "#22543D", "border": "#C6F6D5", "icon": "🟢"},
    "보류": {"bg": "#FAF5FF", "text": "#6B46C1", "border": "#E9D8FD", "icon": "🟣"},
    "취소": {"bg": "#FFF5F5", "text": "#C53030", "border": "#FED7D7", "icon": "⚪"},
}

STATUS_CONFIG_DARK = {
    "등록": {"bg": "#1E3A5F", "text": "#90CDF4", "border": "#2B6CB0", "icon": "🔵"},
    "진행중": {"bg": "#4A2800", "text": "#FBD38D", "border": "#DD6B20", "icon": "🟠"},
    "완료": {"bg": "#1C4532", "text": "#9AE6B4", "border": "#38A169", "icon": "🟢"},
    "보류": {"bg": "#322659", "text": "#D6BCFA", "border": "#805AD5", "icon": "🟣"},
    "취소": {"bg": "#4A1D24", "text": "#FEB2B2", "border": "#E53E3E", "icon": "⚪"},
}

CATEGORY_CONFIG = {
    "일반": {"bg": "#F1F5F9", "text": "#475569", "border": "#CBD5E1"},
    "공문": {"bg": "#EFF6FF", "text": "#1D4ED8", "border": "#BFDBFE"},
    "보고": {"bg": "#F5F3FF", "text": "#6D28D9", "border": "#DDD6FE"},
    "기안": {"bg": "#ECFDF5", "text": "#047857", "border": "#A7F3D0"},
    "결재": {"bg": "#FFFBEB", "text": "#B45309", "border": "#FDE68A"},
    "민원": {"bg": "#FEF2F2", "text": "#B91C1C", "border": "#FECACA"},
    "예산": {"bg": "#F0FDF4", "text": "#15803D", "border": "#BBF7D0"},
    "기타": {"bg": "#F8FAFC", "text": "#64748B", "border": "#E2E8F0"},
}

CATEGORY_CONFIG_DARK = {
    "일반": {"bg": "#2D3748", "text": "#CBD5E0", "border": "#4A5568"},
    "공문": {"bg": "#1E3A8A", "text": "#93C5FD", "border": "#2563EB"},
    "보고": {"bg": "#4C1D95", "text": "#C4B5FD", "border": "#7C3AED"},
    "기안": {"bg": "#064E3B", "text": "#6EE7B7", "border": "#059669"},
    "결재": {"bg": "#78350F", "text": "#FCD34D", "border": "#D97706"},
    "민원": {"bg": "#7F1D1D", "text": "#FCA5A5", "border": "#DC2626"},
    "예산": {"bg": "#14532D", "text": "#86EFAC", "border": "#16A34A"},
    "기타": {"bg": "#1F2937", "text": "#9CA3AF", "border": "#374151"},
}


class RowHoverDelegate(QStyledItemDelegate):
    """테이블 마우스 오버 시 1줄 전체 배경색을 균일하게 칠해주는 델리게이트"""

    def __init__(self, parent=None) -> None:
        super().__init__(parent)
        self.hovered_row: int = -1
        self.hover_color: QColor = QColor("#F8FAFC")

    def paint(self, painter, option, index) -> None:
        if index.row() == self.hovered_row and not (option.state & QStyle.State_Selected):
            painter.fillRect(option.rect, self.hover_color)
        super().paint(painter, option, index)


class PillBadgeWidget(QWidget):
    """알약(Pill) 형태의 깔끔한 뱃지 버튼 위젯 (단일 클릭 드롭다운 메뉴 연결)"""

    def __init__(
        self,
        text: str,
        config: dict,
        on_click,
        on_context=None,
        parent: QWidget | None = None,
    ) -> None:
        super().__init__(parent)
        self.setAttribute(Qt.WA_TranslucentBackground, True)
        self.setStyleSheet("background: transparent;")
        self.on_click = on_click
        self.on_context = on_context
        layout = QHBoxLayout(self)
        layout.setContentsMargins(4, 2, 4, 2)
        layout.setAlignment(Qt.AlignCenter)

        cfg = config.get(text, {"bg": "#F1F5F9", "text": "#475569", "border": "#CBD5E1"})
        bg = cfg.get("bg", "#F1F5F9")
        color = cfg.get("text", "#475569")
        border = cfg.get("border", "#CBD5E1")

        self.btn = QPushButton(f"{text}  ▾")
        self.btn.setCursor(Qt.PointingHandCursor)
        self.btn.setFixedHeight(24)
        self.btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {bg};
                color: {color};
                border: 1px solid {border};
                border-radius: 12px;
                padding: 0 10px;
                font-family: 'Malgun Gothic';
                font-size: 11px;
                font-weight: 600;
            }}
            QPushButton:hover {{
                border-color: {color};
                background-color: #FFFFFF;
            }}
        """)
        self.btn.clicked.connect(lambda: self.on_click(self.btn))
        if self.on_context:
            self.btn.setContextMenuPolicy(Qt.CustomContextMenu)
            self.btn.customContextMenuRequested.connect(lambda p: self.on_context(QCursor.pos()))
        layout.addWidget(self.btn)


class SimpleListManagerDialog(QDialog):
    """항목(분류, 상태 등)을 추가/삭제/관리하는 모달 창"""

    def __init__(
        self,
        parent: QWidget | None,
        title: str,
        items: list[str],
        on_save_callback,
        immutable_item: str = "",
        palette: dict[str, str] | None = None,
    ) -> None:
        super().__init__(parent)
        self.palette = palette or get_dialog_palette(parent)
        self.setWindowTitle(title)
        self.setFixedSize(360, 430)
        self.immutable_item = immutable_item
        # 기본 고정 항목은 항상 맨 앞에 위치
        if self.immutable_item:
            self.items = [self.immutable_item] + [x for x in items if x != self.immutable_item]
        else:
            self.items = list(items)
        self.on_save_callback = on_save_callback

        self.setStyleSheet(f"QDialog {{ background-color: {self.palette['bg']}; color: {self.palette['text']}; }}")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(14, 14, 14, 14)
        layout.setSpacing(10)

        notice = f"<b>{title}</b>"
        muted_c = self.palette.get("muted", "#718096")
        if self.immutable_item:
            notice += f"<br><span style='color: {muted_c}; font-size: 11px;'>※ '{self.immutable_item}'은(는) 기본 고정 항목으로 수정/삭제할 수 없습니다.</span>"
        lbl_notice = QLabel(notice)
        lbl_notice.setStyleSheet(f"color: {self.palette['text']};")
        layout.addWidget(lbl_notice)

        panel = self.palette.get("panel", "#FFFFFF")
        panel_alt = self.palette.get("panel_alt", "#F8FAFC")
        text = self.palette.get("text", "#1F2328")
        line = self.palette.get("line", "#CBD5E0")
        line_soft = self.palette.get("line_soft", line)
        accent = self.palette.get("accent", "#1F7A67")
        accent_soft = self.palette.get("accent_soft", panel_alt)
        btn_text = self.palette.get("button_text", "#FFFFFF")
        danger = self.palette.get("danger", "#E53E3E")

        self.list_widget = QListWidget()
        self.list_widget.setStyleSheet(f"""
            QListWidget {{
                background-color: {panel};
                color: {text};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 4px;
                font-size: 12px;
            }}
            QListWidget::item {{
                padding: 6px;
                border-bottom: 1px solid {line_soft};
            }}
            QListWidget::item:selected {{
                background-color: {accent_soft};
                color: {text};
            }}
        """)
        for item in self.items:
            display_text = f"🔒 {item} (기본 고정)" if item == self.immutable_item else item
            self.list_widget.addItem(display_text)
        layout.addWidget(self.list_widget, 1)

        # 추가 입력창
        add_layout = QHBoxLayout()
        self.new_item_input = QLineEdit()
        self.new_item_input.setPlaceholderText("새 항목 입력...")
        self.new_item_input.setFixedHeight(30)
        self.new_item_input.setStyleSheet(f"""
            QLineEdit {{
                background-color: {panel};
                color: {text};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 4px 8px;
            }}
            QLineEdit:focus {{
                border-color: {accent};
            }}
        """)
        self.new_item_input.returnPressed.connect(self._add_item)
        add_layout.addWidget(self.new_item_input)

        btn_add = QPushButton("추가")
        btn_add.setFixedHeight(30)
        btn_add.setStyleSheet(f"""
            QPushButton {{
                background-color: {accent};
                color: {btn_text};
                font-weight: bold;
                border-radius: 6px;
                padding: 0 14px;
                border: 1px solid {accent};
            }}
            QPushButton:hover {{
                background-color: {_shade(accent, -0.12)};
            }}
        """)
        btn_add.clicked.connect(self._add_item)
        add_layout.addWidget(btn_add)
        layout.addLayout(add_layout)

        # 하단 버튼
        btn_layout = QHBoxLayout()
        btn_del = QPushButton("선택 항목 삭제")
        btn_del.setFixedHeight(30)
        btn_del.setStyleSheet(f"""
            QPushButton {{
                background-color: {panel_alt};
                color: {danger};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 0 10px;
            }}
            QPushButton:hover {{
                background-color: {panel};
                border-color: {danger};
            }}
        """)
        btn_del.clicked.connect(self._delete_item)
        btn_layout.addWidget(btn_del)

        btn_layout.addStretch(1)

        btn_ok = QPushButton("저장 완료")
        btn_ok.setFixedHeight(30)
        btn_ok.setStyleSheet(f"""
            QPushButton {{
                background-color: {panel};
                color: {text};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 0 14px;
                font-weight: bold;
            }}
            QPushButton:hover {{
                background-color: {panel_alt};
            }}
        """)
        btn_ok.clicked.connect(self._save_and_close)
        btn_layout.addWidget(btn_ok)

        layout.addLayout(btn_layout)

    def _add_item(self) -> None:
        text = self.new_item_input.text().strip()
        if not text:
            return
        if text in self.items:
            QMessageBox.information(self, "안내", "이미 존재하는 항목입니다.")
            return
        self.items.append(text)
        self.list_widget.addItem(text)
        self.new_item_input.clear()

    def _delete_item(self) -> None:
        row = self.list_widget.currentRow()
        if row < 0 or row >= len(self.items):
            return
        if self.items[row] == self.immutable_item:
            QMessageBox.warning(self, "삭제 불가", f"기본 항목 '{self.immutable_item}'은(는) 삭제할 수 없습니다.")
            return
        del self.items[row]
        self.list_widget.takeItem(row)

    def _save_and_close(self) -> None:
        if not self.items:
            QMessageBox.warning(self, "경고", "최소 하나 이상의 항목이 필요합니다.")
            return
        self.on_save_callback(self.items)
        self.accept()


class TaskEditDialog(QDialog):
    """업무 전용 등록 및 수정 대화상자 (기안자 자동완성, 분류/상태 관리 포함)"""

    def __init__(
        self,
        parent: QWidget | None,
        repository: EncryptedRepository,
        task: CalendarEntry | None = None,
        known_authors: list[str] | None = None,
        palette: dict[str, str] | None = None,
    ) -> None:
        super().__init__(parent)
        self.palette = palette or get_dialog_palette(parent, repository)
        self.repository = repository
        self.task = task
        self.known_authors = known_authors or []

        is_new = (task is None or task.entry_id is None)
        self.setWindowTitle("새 업무 등록" if is_new else "업무 정보 수정")
        self.resize(620, 560)
        self.setStyleSheet(f"QDialog {{ background-color: {self.palette['bg']}; color: {self.palette['text']}; }}")

        self._init_ui()
        self._load_data()

    def _init_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(12)

        grid = QGridLayout()
        grid.setHorizontalSpacing(12)
        grid.setVerticalSpacing(10)
        grid.setColumnStretch(1, 1)
        grid.setColumnStretch(3, 1)

        panel = self.palette.get("panel", "#FFFFFF")
        panel_alt = self.palette.get("panel_alt", "#F8FAFC")
        text = self.palette.get("text", "#1F2328")
        line = self.palette.get("line", "#CBD5E0")
        accent = self.palette.get("accent", "#1F7A67")
        accent_soft = self.palette.get("accent_soft", panel_alt)
        btn_text = self.palette.get("button_text", "#FFFFFF")

        lbl_style = f"font-size: 12px; font-weight: bold; color: {text};"
        arrow_svg = str(asset_path("chevron_down.svg")).replace("\\", "/")

        combo_style = f"""
            QComboBox {{
                border: 1px solid {line};
                border-radius: 6px;
                padding: 4px 24px 4px 10px;
                font-size: 12px;
                background-color: {panel};
                color: {text};
            }}
            QComboBox:hover, QComboBox:focus {{
                border-color: {accent};
            }}
            QComboBox::drop-down {{
                subcontrol-origin: padding;
                subcontrol-position: top right;
                width: 22px;
                border-left: 1px solid {line};
            }}
            QComboBox::down-arrow {{
                image: url('{arrow_svg}');
                width: 9px;
                height: 6px;
                margin-right: 2px;
            }}
            QComboBox QAbstractItemView {{
                border: 1px solid {line};
                border-radius: 6px;
                background-color: {panel};
                color: {text};
                selection-background-color: {accent_soft};
                selection-color: {text};
                outline: 0px;
                padding: 4px;
            }}
        """

        sub_btn_style = f"""
            QPushButton {{
                background-color: {panel_alt};
                color: {text};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 0 8px;
                font-size: 11px;
            }}
            QPushButton:hover {{
                background-color: {panel};
                border-color: {accent};
            }}
        """

        input_style = f"""
            QLineEdit {{
                background-color: {panel};
                color: {text};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 4px 8px;
                font-size: 12px;
            }}
            QLineEdit:focus {{
                border-color: {accent};
            }}
        """

        # 1. 업무분류 (0,0~0,1) & 등록일자 (0,2~0,3) 한 줄 배치
        lbl_cat = QLabel("업무분류:")
        lbl_cat.setStyleSheet(lbl_style)
        lbl_cat.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        grid.addWidget(lbl_cat, 0, 0)

        cat_layout = QHBoxLayout()
        cat_layout.setContentsMargins(0, 0, 0, 0)
        cat_layout.setSpacing(4)
        self.category_combo = QComboBox()
        self.category_combo.setFixedHeight(30)
        self.category_combo.setStyleSheet(combo_style)
        self._refresh_categories()
        cat_layout.addWidget(self.category_combo, 1)

        btn_cat_mgr = QPushButton("분류 관리")
        btn_cat_mgr.setFixedHeight(30)
        btn_cat_mgr.setStyleSheet(sub_btn_style)
        btn_cat_mgr.clicked.connect(self._open_category_manager)
        cat_layout.addWidget(btn_cat_mgr)
        grid.addLayout(cat_layout, 0, 1)

        lbl_date = QLabel("등록일자:")
        lbl_date.setStyleSheet(lbl_style)
        lbl_date.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        grid.addWidget(lbl_date, 0, 2)

        date_layout = QHBoxLayout()
        date_layout.setContentsMargins(0, 0, 0, 0)
        date_layout.setSpacing(4)
        self.date_edit = SafeDateEdit()
        self.date_edit.setCalendarPopup(True)
        self.date_edit.setDisplayFormat("yyyy-MM-dd")
        self.date_edit.setFixedHeight(30)
        self.date_edit.setDate(date.today())
        self.date_edit.setStyleSheet(
            f"QDateEdit {{ background: {panel}; color: {text}; border: 1px solid {line}; border-radius: 6px; padding: 2px 24px 2px 8px; font-size: 13px; }}"
            f"QDateEdit:hover, QDateEdit:focus {{ border-color: {accent}; }}"
            f"QDateEdit::drop-down {{ subcontrol-origin: padding; subcontrol-position: top right; width: 22px; border-left: 1px solid {line}; }}"
            f"QDateEdit::down-arrow {{ image: url('{arrow_svg}'); width: 9px; height: 6px; margin-right: 2px; }}"
        )
        date_layout.addWidget(self.date_edit, 1)

        btn_date_today = QPushButton("오늘")
        btn_date_today.setFixedHeight(30)
        btn_date_today.setStyleSheet(sub_btn_style)
        btn_date_today.clicked.connect(lambda: self.date_edit.setDate(date.today()))
        date_layout.addWidget(btn_date_today)
        grid.addLayout(date_layout, 0, 3)

        # 2. 업무 제목 (1,0~1,3 전체 가로 확장)
        lbl_title = QLabel("업무제목:")
        lbl_title.setStyleSheet(lbl_style)
        lbl_title.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        grid.addWidget(lbl_title, 1, 0)

        self.title_input = QLineEdit()
        self.title_input.setPlaceholderText("온나라 결재문서 제목 또는 업무 명칭")
        self.title_input.setFixedHeight(30)
        self.title_input.setStyleSheet(f"""
            QLineEdit {{
                background-color: {panel};
                color: {text};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 4px 8px;
                font-size: 13px;
                font-weight: bold;
            }}
            QLineEdit:focus {{
                border-color: {accent};
            }}
        """)
        grid.addWidget(self.title_input, 1, 1, 1, 3)

        # 3. 기안자 (2,0~2,1) & 처리상태 (2,2~2,3) 한 줄 배치
        lbl_author = QLabel("기안자:")
        lbl_author.setStyleSheet(lbl_style)
        lbl_author.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        grid.addWidget(lbl_author, 2, 0)

        self.author_input = QLineEdit()
        self.author_input.setPlaceholderText("기안자 성명/부서 (예: 홍길동)")
        self.author_input.setFixedHeight(30)
        self.author_input.setStyleSheet(input_style)
        self._setup_author_completer()
        grid.addWidget(self.author_input, 2, 1)

        lbl_status = QLabel("처리상태:")
        lbl_status.setStyleSheet(lbl_style)
        lbl_status.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        grid.addWidget(lbl_status, 2, 2)

        status_layout = QHBoxLayout()
        status_layout.setContentsMargins(0, 0, 0, 0)
        status_layout.setSpacing(4)
        self.status_combo = QComboBox()
        self.status_combo.setFixedHeight(30)
        self.status_combo.setStyleSheet(combo_style)
        self._refresh_statuses()
        status_layout.addWidget(self.status_combo, 1)

        btn_status_mgr = QPushButton("상태 관리")
        btn_status_mgr.setFixedHeight(30)
        btn_status_mgr.setStyleSheet(sub_btn_style)
        btn_status_mgr.clicked.connect(self._open_status_manager)
        status_layout.addWidget(btn_status_mgr)
        grid.addLayout(status_layout, 2, 3)

        # 4. 비고 / 세부내용 (남은 수직 공간을 모두 차지하여 시원하게 확장)
        lbl_desc = QLabel("비고/내용:")
        lbl_desc.setStyleSheet(lbl_style)
        lbl_desc.setAlignment(Qt.AlignRight | Qt.AlignTop)
        grid.addWidget(lbl_desc, 3, 0)

        self.desc_input = QTextEdit()
        self.desc_input.setPlaceholderText("결재 요지, 업무 세부내용, 지시사항 등")
        self.desc_input.setMinimumHeight(140)
        self.desc_input.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.desc_input.setStyleSheet(f"""
            QTextEdit {{
                background-color: {panel};
                color: {text};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 8px;
                font-size: 12px;
            }}
            QTextEdit:focus {{
                border-color: {accent};
            }}
        """)
        grid.addWidget(self.desc_input, 3, 1, 1, 3)
        grid.setRowStretch(3, 1)

        # 5. 출처 URL (4,0~4,3)
        lbl_url = QLabel("출처 URL:")
        lbl_url.setStyleSheet(lbl_style)
        lbl_url.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        grid.addWidget(lbl_url, 4, 0)

        self.url_input = QLineEdit()
        self.url_input.setPlaceholderText("온나라 문서 링크 또는 웹사이트 URL")
        self.url_input.setFixedHeight(28)
        self.url_input.setStyleSheet(input_style)
        grid.addWidget(self.url_input, 4, 1, 1, 3)

        layout.addLayout(grid, 1)

        # 하단 확인 / 취소 버튼
        btn_box = QHBoxLayout()
        btn_box.setSpacing(6)
        btn_box.addStretch(1)

        btn_cancel = QPushButton("취소")
        btn_cancel.setFixedHeight(28)
        btn_cancel.setStyleSheet(f"""
            QPushButton {{
                background-color: {panel};
                color: {text};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 0 14px;
                font-size: 12px;
            }}
            QPushButton:hover {{
                background-color: {panel_alt};
            }}
        """)
        btn_cancel.clicked.connect(self.reject)
        btn_box.addWidget(btn_cancel)

        btn_save = QPushButton("저장")
        btn_save.setFixedHeight(28)
        btn_save.setStyleSheet(f"""
            QPushButton {{
                background-color: {accent};
                color: {btn_text};
                border: 1px solid {accent};
                border-radius: 6px;
                padding: 0 16px;
                font-size: 12px;
                font-weight: 600;
            }}
            QPushButton:hover {{
                background-color: {_shade(accent, -0.12)};
                border-color: {_shade(accent, -0.12)};
            }}
            QPushButton:pressed {{
                background-color: {_shade(accent, -0.20)};
                border-color: {_shade(accent, -0.20)};
            }}
        """)
        btn_save.clicked.connect(self._on_save)
        btn_box.addWidget(btn_save)

        layout.addLayout(btn_box)

    def _setup_author_completer(self) -> None:
        """기존 기안자 목록으로 자동완성 세팅"""
        clean_authors = sorted(list({a.strip() for a in self.known_authors if a and a.strip()}))
        completer = QCompleter(clean_authors, self)
        completer.setCaseSensitivity(Qt.CaseInsensitive)
        completer.setFilterMode(Qt.MatchContains)  # 중간 글자 입력해도 매칭
        self.author_input.setCompleter(completer)

    def _refresh_categories(self) -> None:
        curr = self.category_combo.currentText()
        cats = get_task_categories(self.repository)
        self.category_combo.clear()
        self.category_combo.addItems(cats)
        if curr and curr in cats:
            self.category_combo.setCurrentText(curr)

    def _refresh_statuses(self) -> None:
        curr = self.status_combo.currentText()
        statuses = get_task_statuses(self.repository)
        self.status_combo.clear()
        self.status_combo.addItems(statuses)
        if curr and curr in statuses:
            self.status_combo.setCurrentText(curr)

    def _open_category_manager(self) -> None:
        cats = get_task_categories(self.repository)
        dlg = SimpleListManagerDialog(
            self,
            "업무 분류 관리",
            cats,
            lambda new_cats: save_task_categories(self.repository, new_cats),
            immutable_item="일반",
            palette=self.palette,
        )
        if dlg.exec() == QDialog.Accepted:
            self._refresh_categories()

    def _open_status_manager(self) -> None:
        statuses = get_task_statuses(self.repository)
        dlg = SimpleListManagerDialog(
            self,
            "처리 상태 관리",
            statuses,
            lambda new_statuses: save_task_statuses(self.repository, new_statuses),
            immutable_item="등록",
            palette=self.palette,
        )
        if dlg.exec() == QDialog.Accepted:
            self._refresh_statuses()

    def _load_data(self) -> None:
        if not self.task:
            self.category_combo.setCurrentText("일반")
            self.status_combo.setCurrentText("등록")
            return

        load_date = self.task.day or self.task.start_date or (self.task.created_at.date() if self.task.created_at else date.today())
        self.date_edit.setDate(load_date)
        self.title_input.setText(self.task.title or "")
        self.author_input.setText(self.task.assignee or "")

        cat = self.task.memo_group or "일반"
        idx = self.category_combo.findText(cat)
        if idx >= 0:
            self.category_combo.setCurrentIndex(idx)
        else:
            self.category_combo.addItem(cat)
            self.category_combo.setCurrentText(cat)

        st = self.task.status or "등록"
        s_idx = self.status_combo.findText(st)
        if s_idx >= 0:
            self.status_combo.setCurrentIndex(s_idx)
        else:
            self.status_combo.addItem(st)
            self.status_combo.setCurrentText(st)

        # 출처 URL과 설명 분리 파싱
        desc = self.task.description or ""
        url = ""
        if "📎 출처:" in desc:
            parts = desc.rsplit("📎 출처:", 1)
            desc = parts[0].strip()
            url = parts[1].strip()
        elif "\n출처:" in desc:
            parts = desc.rsplit("\n출처:", 1)
            desc = parts[0].strip()
            url = parts[1].strip()
        elif "출처:" in desc:
            parts = desc.rsplit("출처:", 1)
            desc = parts[0].strip()
            url = parts[1].strip()

        self.desc_input.setPlainText(desc)
        self.url_input.setText(url)

    def _on_save(self) -> None:
        title = self.title_input.text().strip()
        if not title:
            QMessageBox.warning(self, "입력 확인", "업무 제목을 입력해 주세요.")
            self.title_input.setFocus()
            return

        target_date = self.date_edit.date().toPython()
        category = self.category_combo.currentText().strip() or "일반"
        author = self.author_input.text().strip()
        status = self.status_combo.currentText().strip() or "등록"
        desc = self.desc_input.toPlainText().strip()
        url = self.url_input.text().strip()

        full_desc = desc
        if url:
            full_desc = f"{desc}\n\n출처: {url}".strip()

        now = datetime.now()
        if not self.task:
            self.task = CalendarEntry(
                entry_type=EntryType.TASK,
                title=title,
                description=full_desc,
                day=target_date,
                start_date=target_date,
                assignee=author,
                memo_group=category,
                status=status,
                created_at=datetime.combine(target_date, now.time()),
            )
        else:
            self.task.entry_type = EntryType.TASK
            self.task.title = title
            self.task.description = full_desc
            self.task.day = target_date
            self.task.start_date = target_date
            if self.task.created_at:
                self.task.created_at = datetime.combine(target_date, self.task.created_at.time())
            else:
                self.task.created_at = datetime.combine(target_date, now.time())
            self.task.assignee = author
            self.task.memo_group = category
            self.task.status = status
            self.task.updated_at = now

        self.repository.upsert_entry(self.task)
        self.repository.save()
        self.accept()


class TaskManagerDialog(QDialog):
    """
    온나라 전자결재 및 웹 업무를 엑셀 표 스타일로 한눈에 조회/관리/검색하고
    엑셀(.xlsx)로 내보내는 업무 관리 대시보드 창
    """

    def __init__(self, parent: QWidget | None, repository: EncryptedRepository, main_window=None) -> None:
        super().__init__(parent)
        self.repository = repository
        self.main_window = main_window
        self.palette = get_dialog_palette(parent, repository, main_window)

        self.setWindowTitle("업무")
        self.setWindowFlags(self.windowFlags() | Qt.WindowMaximizeButtonHint | Qt.WindowMinimizeButtonHint)
        self.resize(1100, 700)
        self.setMinimumSize(850, 520)
        self.setAttribute(Qt.WA_StyledBackground, True)

        self._all_tasks: list[CalendarEntry] = []
        self._filtered_tasks: list[CalendarEntry] = []
        self._displayed_tasks: list[CalendarEntry] = []
        self._current_tab = "all"
        self._current_page = 1
        self._current_date_anchor = None

        self._init_ui()
        self.reload_tasks()

    def _init_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 14, 16, 14)
        layout.setSpacing(10)

        # 1. 상단 헤더: 메인 액션 버튼 [등록] [설정] [엑셀저장] (같은 줄, 우측 정렬)
        header_layout = QHBoxLayout()
        header_layout.setContentsMargins(0, 0, 0, 0)
        header_layout.setSpacing(8)
        header_layout.addStretch(1)

        # 신규 등록 버튼 (검색 버튼과 동일한 크기(세로 28px) 및 테마 강조색 스타일)
        self.btn_add = QPushButton("등록")
        self.btn_add.setFixedHeight(28)
        self.btn_add.clicked.connect(self._on_add_task)
        header_layout.addWidget(self.btn_add)

        # 분류/상태 설정 버튼 (세로 높이 28px)
        self.btn_manage_meta = QPushButton("설정")
        self.btn_manage_meta.setFixedHeight(28)
        self.btn_manage_meta.clicked.connect(self._open_meta_manager_menu)
        header_layout.addWidget(self.btn_manage_meta)

        # 엑셀저장 버튼 (세로 높이 28px)
        self.btn_export = QPushButton("엑셀저장")
        self.btn_export.setFixedHeight(28)
        self.btn_export.clicked.connect(self._export_to_excel)
        header_layout.addWidget(self.btn_export)

        layout.addLayout(header_layout)

        # 2. 검색 및 필터 툴바 카드 (분류, 상태, 기간(시작일~종료일), 제목+내용 검색 순서)
        self.toolbar_card = QFrame()
        self.toolbar_card.setObjectName("toolbar_card")
        tb_layout = QHBoxLayout(self.toolbar_card)
        tb_layout.setContentsMargins(12, 8, 12, 8)
        tb_layout.setSpacing(10)

        # (1) 분류
        self.category_combo = QComboBox()
        self.category_combo.setFixedHeight(32)
        self.category_combo.setFixedWidth(85)
        if self.category_combo.view():
            self.category_combo.view().setMinimumWidth(100)
        self._refresh_filter_categories()
        self.category_combo.currentIndexChanged.connect(self._apply_filters)
        tb_layout.addWidget(self.category_combo)

        # (2) 상태
        self.status_combo = QComboBox()
        self.status_combo.setFixedHeight(32)
        self.status_combo.setFixedWidth(85)
        if self.status_combo.view():
            self.status_combo.view().setMinimumWidth(90)
        self._refresh_filter_statuses()
        self.status_combo.currentIndexChanged.connect(self._apply_filters)
        tb_layout.addWidget(self.status_combo)

        # (3) 기간(시작일~종료일)
        self.period_combo = QComboBox()
        self.period_combo.setFixedHeight(32)
        self.period_combo.setFixedWidth(100)
        if self.period_combo.view():
            self.period_combo.view().setMinimumWidth(110)
        self.period_combo.addItems(["기간", "오늘", "이번 주", "이번 달", "최근 3개월", "올해", "직접 지정"])
        self.period_combo.currentIndexChanged.connect(self._on_period_combo_changed)
        tb_layout.addWidget(self.period_combo)

        self.start_date_edit = SafeDateEdit()
        self.start_date_edit.setCalendarPopup(True)
        self.start_date_edit.setDisplayFormat("yyyy-MM-dd")
        self.start_date_edit.setFixedHeight(32)
        self.start_date_edit.setFixedWidth(125)
        self.start_date_edit.setDate(date.today() - timedelta(days=30))
        self.start_date_edit.setEnabled(True)
        self.start_date_edit.dateChanged.connect(self._on_custom_date_changed)
        tb_layout.addWidget(self.start_date_edit)

        self.lbl_tilde = QLabel("~")
        tb_layout.addWidget(self.lbl_tilde)

        self.end_date_edit = SafeDateEdit()
        self.end_date_edit.setCalendarPopup(True)
        self.end_date_edit.setDisplayFormat("yyyy-MM-dd")
        self.end_date_edit.setFixedHeight(32)
        self.end_date_edit.setFixedWidth(125)
        self.end_date_edit.setDate(date.today())
        self.end_date_edit.setEnabled(True)
        self.end_date_edit.dateChanged.connect(self._on_custom_date_changed)
        tb_layout.addWidget(self.end_date_edit)

        # (4) 제목+내용 검색 (남은 가로 공간 확장)
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("🔍  제목+내용 검색...")
        self.search_input.setFixedHeight(32)
        self.search_input.setClearButtonEnabled(True)
        self.search_input.textChanged.connect(self._apply_filters)
        tb_layout.addWidget(self.search_input, 1)

        layout.addWidget(self.toolbar_card)

        # 3. 하단 리스트 영역: 탭 바 + 테이블 묶음 (탭과 표 사이 줄간격 0)
        table_container = QVBoxLayout()
        table_container.setContentsMargins(0, 0, 0, 0)
        table_container.setSpacing(0)

        self.tabs_widget = QWidget()
        self.tabs_widget.setFixedHeight(28)
        tabs_bar = QHBoxLayout(self.tabs_widget)
        tabs_bar.setContentsMargins(0, 0, 0, 0)
        tabs_bar.setSpacing(0)

        self.btn_tab_all = QPushButton("전체")
        self.btn_tab_reg = QPushButton("등록")
        self.btn_tab_prog = QPushButton("진행")
        self.btn_tab_done = QPushButton("완료")

        for btn in (self.btn_tab_all, self.btn_tab_reg, self.btn_tab_prog, self.btn_tab_done):
            btn.setCursor(Qt.PointingHandCursor)

        self.btn_tab_all.clicked.connect(lambda: self._set_active_tab("all"))
        self.btn_tab_reg.clicked.connect(lambda: self._set_active_tab("reg"))
        self.btn_tab_prog.clicked.connect(lambda: self._set_active_tab("prog"))
        self.btn_tab_done.clicked.connect(lambda: self._set_active_tab("done"))

        tabs_bar.addWidget(self.btn_tab_all, 0, Qt.AlignTop)
        tabs_bar.addWidget(self.btn_tab_reg, 0, Qt.AlignTop)
        tabs_bar.addWidget(self.btn_tab_prog, 0, Qt.AlignTop)
        tabs_bar.addWidget(self.btn_tab_done, 0, Qt.AlignTop)
        tabs_bar.addStretch(1)

        # 오른쪽 끝: 페이징 네비게이션 ([◀] [정보] [▶]) + 페이징 단위 선택 셀렉트(QComboBox)
        self.paging_nav_widget = QWidget()
        self.paging_nav_widget.setFixedHeight(26)
        nav_layout = QHBoxLayout(self.paging_nav_widget)
        nav_layout.setContentsMargins(0, 0, 4, 0)
        nav_layout.setSpacing(4)
        nav_layout.setSizeConstraint(QLayout.SetFixedSize)

        self.btn_page_prev = QPushButton("◀")
        self.btn_page_prev.setFixedSize(22, 22)
        self.btn_page_prev.setCursor(Qt.PointingHandCursor)
        self.btn_page_prev.clicked.connect(self._on_page_prev)

        self.lbl_page_info = QLabel("")
        self.lbl_page_info.setAlignment(Qt.AlignCenter)

        self.btn_page_next = QPushButton("▶")
        self.btn_page_next.setFixedSize(22, 22)
        self.btn_page_next.setCursor(Qt.PointingHandCursor)
        self.btn_page_next.clicked.connect(self._on_page_next)

        nav_layout.addWidget(self.btn_page_prev)
        nav_layout.addWidget(self.lbl_page_info)
        nav_layout.addWidget(self.btn_page_next)
        tabs_bar.addWidget(self.paging_nav_widget, 0, Qt.AlignVCenter)

        self.paging_combo = QComboBox()
        self.paging_combo.setFixedHeight(24)
        self.paging_combo.setFixedWidth(110)
        self.paging_combo.addItems([
            "전체",
            "10개씩",
            "20개씩",
            "50개씩",
            "일별",
            "주간별",
            "월별",
        ])
        saved_paging = self.repository.get_setting("task_paging_mode") or "전체"
        p_idx = self.paging_combo.findText(saved_paging)
        if p_idx >= 0:
            self.paging_combo.setCurrentIndex(p_idx)
        else:
            self.paging_combo.setCurrentIndex(0)
        self.paging_combo.currentIndexChanged.connect(self._on_paging_combo_changed)
        tabs_bar.addWidget(self.paging_combo, 0, Qt.AlignVCenter)

        table_container.addWidget(self.tabs_widget)

        # 4. 하단 리스트 테이블 (QTableWidget) - 7개 컬럼 구성
        self.table = QTableWidget()
        self.table.setColumnCount(7)
        self.table.verticalHeader().setVisible(False)
        self.table.setAlternatingRowColors(False)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setSelectionMode(QTableWidget.SingleSelection)
        self.table.setShowGrid(True)

        self.HEADER_COLUMNS = ["등록일자", "분류", "업무 제목", "기안자 / 작성자", "상태", "비고 / 세부내용", "편집"]
        self._sort_column: int = 0
        self._sort_order: Qt.SortOrder = Qt.DescendingOrder
        self._update_header_labels()

        header = self.table.horizontalHeader()
        header.setSectionsClickable(True)
        header.setSortIndicatorShown(False)
        header.sectionClicked.connect(self._on_header_clicked)

        # 헤더 좌우 드래그로 모든 셀/컬럼 크기 조절 가능하도록 Interactive 모드로 설정
        header.setSectionResizeMode(QHeaderView.Interactive)
        header.setSectionResizeMode(6, QHeaderView.Fixed)  # 편집 아이콘 열은 55px 고정

        self._default_col_widths = [110, 95, 280, 130, 95, 270, 55]
        self._restore_column_widths()

        self._save_columns_timer = QTimer(self)
        self._save_columns_timer.setSingleShot(True)
        self._save_columns_timer.setInterval(400)
        self._save_columns_timer.timeout.connect(self._save_column_widths)

        header.sectionResized.connect(self._on_section_resized)

        self.table.verticalHeader().setDefaultSectionSize(38)

        # 우클릭 컨텍스트 메뉴 설정
        self.table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self._show_context_menu)

        # 테이블 내 직접 수정 리스너
        self.table.itemChanged.connect(self._on_table_item_changed)

        # 마우스 오버 시 1줄 전체 배경색 하이라이트를 위한 마우스 추적 및 델리게이트
        self.hover_delegate = RowHoverDelegate(self.table)
        self.table.setItemDelegate(self.hover_delegate)
        self.table.setMouseTracking(True)
        self.table.viewport().setMouseTracking(True)
        self.table.viewport().installEventFilter(self)

        # 테이블을 table_container에 추가하고 메인 레이아웃에 배치 (탭과의 줄간격 0)
        table_container.addWidget(self.table, 1)
        layout.addLayout(table_container, 1)
        self.tabs_widget.raise_()

        # 스킨/테마 스타일 일괄 적용
        self._apply_dialog_styles()

        # 기본적으로 전체 탭 활성화
        self._set_active_tab("all")

    def _apply_dialog_styles(self) -> None:
        """현재 self.palette를 바탕으로 창 배경, 버튼, 툴바, 탭, 테이블 스타일 일괄 갱신"""
        self.setStyleSheet(f"QDialog {{ background-color: {self.palette['bg']}; color: {self.palette['text']}; }}")
        self._apply_header_button_styles()
        self._apply_toolbar_styles()
        self._apply_table_styles()
        self._update_tab_styles()
        self._apply_paging_styles()

    def _apply_paging_styles(self) -> None:
        """페이징 셀렉트 및 네비게이션 버튼 스타일 갱신"""
        if not hasattr(self, "paging_combo") or not hasattr(self, "btn_page_prev"):
            return
        panel = self.palette.get("panel", "#FFFFFF")
        panel_alt = self.palette.get("panel_alt", "#F8FAFC")
        text = self.palette.get("text", "#1F2328")
        line = self.palette.get("line", "#CBD5E0")
        accent = self.palette.get("accent", "#1F7A67")
        arrow_svg = str(asset_path("chevron_down.svg")).replace("\\", "/")

        self.paging_combo.setStyleSheet(f"""
            QComboBox {{
                border: 1px solid {line};
                border-radius: 4px;
                padding: 1px 16px 1px 6px;
                font-size: 11px;
                font-weight: 500;
                background-color: {panel};
                color: {text};
            }}
            QComboBox:hover {{
                border-color: {accent};
            }}
            QComboBox::drop-down {{
                subcontrol-origin: padding;
                subcontrol-position: top right;
                width: 14px;
                border-left: none;
            }}
            QComboBox::down-arrow {{
                image: url("{arrow_svg}");
                width: 8px;
                height: 5px;
                margin-right: 3px;
            }}
            QComboBox QAbstractItemView {{
                border: 1px solid {line};
                border-radius: 4px;
                background-color: {panel};
                color: {text};
                padding: 2px;
                outline: 0px;
            }}
        """)

        btn_style = f"""
            QPushButton {{
                border: 1px solid {line};
                border-radius: 4px;
                background-color: {panel};
                color: {text};
                font-size: 10px;
                font-weight: bold;
                padding: 0px;
            }}
            QPushButton:hover {{
                background-color: {panel_alt};
                border-color: {accent};
            }}
            QPushButton:disabled {{
                color: {line};
                background-color: transparent;
                border-color: {line};
            }}
        """
        self.btn_page_prev.setStyleSheet(btn_style)
        self.btn_page_next.setStyleSheet(btn_style)
        self.lbl_page_info.setStyleSheet(f"font-size: 11px; font-weight: 600; color: {text};")

    def _apply_header_button_styles(self) -> None:
        """상단 [등록], [설정], [엑셀저장] 버튼 스킨 스타일 갱신"""
        accent = self.palette.get("accent", "#1F7A67")
        btn_text = self.palette.get("button_text", "#FFFFFF")
        accent_hover = _shade(accent, -0.12)
        accent_pressed = _shade(accent, -0.20)

        # 신규 등록 버튼: 검색 버튼과 동일한 테마 강조색(accent)
        self.btn_add.setStyleSheet(f"""
            QPushButton {{
                background-color: {accent};
                color: {btn_text};
                border: 1px solid {accent};
                border-radius: 6px;
                padding: 5px 10px;
                font-size: 12px;
                font-weight: 600;
            }}
            QPushButton:hover {{
                background-color: {accent_hover};
                border-color: {accent_hover};
            }}
            QPushButton:pressed {{
                background-color: {accent_pressed};
                border-color: {accent_pressed};
            }}
        """)

        # [설정], [엑셀저장] 버튼: 패널 배경색과 텍스트색
        panel = self.palette.get("panel", "#FFFFFF")
        text = self.palette.get("text", "#1F2328")
        line = self.palette.get("line", "#CBD5E0")
        panel_alt = self.palette.get("panel_alt", "#F6F8FB")
        muted = self.palette.get("muted", line)

        sub_btn_style = f"""
            QPushButton {{
                background-color: {panel};
                color: {text};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 5px 10px;
                font-size: 12px;
            }}
            QPushButton:hover {{
                background-color: {panel_alt};
                border-color: {muted};
            }}
        """
        self.btn_manage_meta.setStyleSheet(sub_btn_style)
        self.btn_export.setStyleSheet(sub_btn_style)

    def _apply_toolbar_styles(self) -> None:
        """검색/필터 카드 및 하위 입력 위젯 스킨 스타일 갱신"""
        panel = self.palette.get("panel", "#FFFFFF")
        panel_alt = self.palette.get("panel_alt", "#F8FAFC")
        text = self.palette.get("text", "#1F2328")
        muted = self.palette.get("muted", "#718096")
        line = self.palette.get("line", "#CBD5E0")
        line_soft = self.palette.get("line_soft", line)
        accent = self.palette.get("accent", "#1F7A67")
        accent_soft = self.palette.get("accent_soft", panel_alt)

        self.toolbar_card.setStyleSheet(f"""
            QFrame#toolbar_card {{
                background-color: {panel};
                border: 1px solid {line_soft};
                border-radius: 8px;
            }}
            QLabel {{
                border: none;
                background-color: transparent;
                color: {text};
            }}
        """)

        arrow_svg = str(asset_path("chevron_down.svg")).replace("\\", "/")

        combo_style = f"""
            QComboBox {{
                border: 1px solid {line};
                border-radius: 6px;
                padding: 4px 18px 4px 8px;
                font-size: 12px;
                background-color: {panel};
                color: {text};
            }}
            QComboBox:hover, QComboBox:focus {{
                border-color: {accent};
            }}
            QComboBox::drop-down {{
                subcontrol-origin: padding;
                subcontrol-position: top right;
                width: 16px;
                border-left: none;
            }}
            QComboBox::down-arrow {{
                image: url("{arrow_svg}");
                width: 9px;
                height: 6px;
                margin-right: 4px;
            }}
            QComboBox QAbstractItemView {{
                border: 1px solid {line};
                border-radius: 6px;
                background-color: {panel};
                color: {text};
                selection-background-color: {accent_soft};
                selection-color: {text};
                outline: 0px;
                padding: 4px;
            }}
        """
        self.category_combo.setStyleSheet(combo_style)
        self.status_combo.setStyleSheet(combo_style)
        self.period_combo.setStyleSheet(combo_style)

        date_edit_style = f"""
            QDateEdit {{
                border: 1px solid {line};
                border-radius: 6px;
                padding: 4px 18px 4px 8px;
                font-size: 12px;
                background-color: {panel};
                color: {text};
            }}
            QDateEdit:hover, QDateEdit:focus {{
                border-color: {accent};
            }}
            QDateEdit::drop-down {{
                subcontrol-origin: padding;
                subcontrol-position: top right;
                width: 16px;
                border-left: none;
            }}
            QDateEdit::down-arrow {{
                image: url("{arrow_svg}");
                width: 9px;
                height: 6px;
                margin-right: 4px;
            }}
        """
        self.start_date_edit.setStyleSheet(date_edit_style)
        self.end_date_edit.setStyleSheet(date_edit_style)

        if hasattr(self, "lbl_tilde"):
            self.lbl_tilde.setStyleSheet(f"border: none; background: transparent; color: {muted}; font-weight: bold;")

        self.search_input.setStyleSheet(f"""
            QLineEdit {{
                border: 1px solid {line};
                border-radius: 6px;
                padding: 4px 10px;
                font-size: 12px;
                background-color: {panel_alt};
                color: {text};
            }}
            QLineEdit:focus {{
                border-color: {accent};
                background-color: {panel};
            }}
        """)

    def _apply_table_styles(self) -> None:
        """하단 리스트 테이블 스킨 스타일 갱신"""
        panel = self.palette.get("panel", "#FFFFFF")
        panel_alt = self.palette.get("panel_alt", "#F8FAFC")
        panel_alt_hover = _shade(panel_alt, -0.06)
        text = self.palette.get("text", "#1F2328")
        muted = self.palette.get("muted", "#667085")
        line = self.palette.get("line", "#CBD5E0")
        line_soft = self.palette.get("line_soft", line)
        accent_soft = self.palette.get("accent_soft", "#EEF2FF")

        self.table.setStyleSheet(f"""
            QTableWidget {{
                background-color: {panel};
                border: 1px solid {line};
                border-top: 1px solid {line};
                border-radius: 0px 6px 6px 6px;
                gridline-color: {line_soft};
                font-size: 13px;
                outline: 0px;
                selection-background-color: {accent_soft};
                selection-color: {text};
            }}
            QTableWidget::item {{
                border: none;
                padding: 4px 8px;
                outline: 0px;
            }}
            QTableWidget::item:focus {{
                border: none;
                outline: 0px;
            }}
            QTableWidget::item:selected {{
                background-color: {accent_soft};
                color: {text};
                border: none;
                outline: 0px;
            }}
            QHeaderView::section {{
                background-color: {panel_alt};
                color: {muted};
                font-size: 12px;
                font-weight: bold;
                padding: 8px 8px;
                border: none;
                border-bottom: 1px solid {line};
                border-right: 1px solid {line_soft};
            }}
        """)

    def _set_hovered_row(self, row: int) -> None:
        """마우스 오버 시 1줄 전체 배경색 하이라이트"""
        if not hasattr(self, "hover_delegate"):
            return
        if row == self.hover_delegate.hovered_row:
            return
        self.hover_delegate.hovered_row = row
        self.table.viewport().update()

    def eventFilter(self, obj: QObject, event: QEvent) -> bool:
        """테이블 및 셀 위젯 마우스 오버 시 1줄 전체 호버 배경색 동기화 처리"""
        if event.type() in (QEvent.MouseMove, QEvent.Enter):
            row_val = obj.property("row_idx")
            if row_val is not None:
                self._set_hovered_row(int(row_val))
            elif obj == self.table.viewport():
                if hasattr(event, "position"):
                    y = int(event.position().y())
                elif hasattr(event, "pos"):
                    y = event.pos().y()
                else:
                    y = self.table.viewport().mapFromGlobal(QCursor.pos()).y()
                r = self.table.rowAt(y)
                self._set_hovered_row(r)
        elif event.type() == QEvent.Leave:
            vp_pos = self.table.viewport().mapFromGlobal(QCursor.pos())
            if not self.table.viewport().rect().contains(vp_pos):
                self._set_hovered_row(-1)
        return super().eventFilter(obj, event)

    def apply_palette(self, palette: dict[str, str]) -> None:
        """스킨/테마 변경 시 팔레트 갱신 및 모든 UI 요소 실시간 리스타일링"""
        self.palette = palette
        if hasattr(self, "hover_delegate"):
            self.hover_delegate.hover_color = QColor(self.palette.get("panel_alt", "#F8FAFC"))
        self._apply_dialog_styles()
        self._render_table()

    def _menu_stylesheet(self) -> str:
        """컨텍스트 메뉴 및 뱃지 드롭다운 메뉴 스킨 스타일 반환"""
        panel = self.palette.get("panel", "#FFFFFF")
        text = self.palette.get("text", "#1F2328")
        line = self.palette.get("line", "#CBD5E0")
        line_soft = self.palette.get("line_soft", line)
        accent_soft = self.palette.get("accent_soft", self.palette.get("panel_alt", "#EDF2F7"))

        return f"""
            QMenu {{
                background-color: {panel};
                border: 1px solid {line};
                border-radius: 6px;
                padding: 4px;
                font-family: 'Malgun Gothic';
                font-size: 12px;
            }}
            QMenu::item {{
                padding: 6px 20px 6px 12px;
                border-radius: 4px;
                color: {text};
            }}
            QMenu::item:selected {{
                background-color: {accent_soft};
                color: {text};
            }}
            QMenu::separator {{
                height: 1px;
                background-color: {line_soft};
                margin: 4px 6px;
            }}
        """

    def _refresh_filter_categories(self) -> None:
        curr = self.category_combo.currentText()
        cats = ["분류"] + get_task_categories(self.repository)
        self.category_combo.clear()
        self.category_combo.addItems(cats)
        if curr and curr in cats:
            self.category_combo.setCurrentText(curr)

    def _refresh_filter_statuses(self) -> None:
        curr = self.status_combo.currentText() if hasattr(self, "status_combo") else ""
        statuses = ["상태"] + get_task_statuses(self.repository)
        self.status_combo.blockSignals(True)
        self.status_combo.clear()
        self.status_combo.addItems(statuses)
        if curr and curr in statuses:
            self.status_combo.setCurrentText(curr)
        else:
            self.status_combo.setCurrentIndex(0)
        self.status_combo.blockSignals(False)

    def _update_tab_styles(self) -> None:
        """상태 탭 스타일 갱신: 스킨 팔레트 반영 및 활성 탭-테이블 일체형 연결"""
        tab_buttons = [
            ("all", self.btn_tab_all),
            ("reg", self.btn_tab_reg),
            ("prog", self.btn_tab_prog),
            ("done", self.btn_tab_done),
        ]
        visible_tabs = [(key, btn) for key, btn in tab_buttons if not btn.isHidden()]
        if not visible_tabs:
            return

        panel = self.palette.get("panel", "#FFFFFF")
        panel_alt = self.palette.get("panel_alt", "#EDF2F7")
        text = self.palette.get("text", "#1F2328")
        muted = self.palette.get("muted", "#4A5568")
        line = self.palette.get("line", "#CBD5E0")
        panel_alt_hover = _shade(panel_alt, -0.06)

        for idx, (key, btn) in enumerate(visible_tabs):
            is_active = (key == self._current_tab)
            is_first = (idx == 0)
            is_last = (idx == len(visible_tabs) - 1)

            # 첫 번째 탭만 왼쪽 모서리 둥글게, 마지막 탭만 오른쪽 모서리 둥글게
            tl_radius = "6px" if is_first else "0px"
            tr_radius = "6px" if is_last else "0px"
            # 첫 번째 탭만 왼쪽 테두리를 그리고, 이후 탭은 왼쪽 테두리를 생략하여 1px 경계 공유
            border_left = f"1px solid {line}" if is_first else "none"

            if is_active:
                btn.setFixedHeight(29)
                btn.setStyleSheet(f"""
                    QPushButton {{
                        background-color: {panel};
                        color: {text};
                        font-weight: bold;
                        border-top: 1px solid {line};
                        border-left: {border_left};
                        border-right: 1px solid {line};
                        border-bottom: 1px solid {panel};
                        border-top-left-radius: {tl_radius};
                        border-top-right-radius: {tr_radius};
                        border-bottom-left-radius: 0px;
                        border-bottom-right-radius: 0px;
                        padding: 4px 16px;
                        font-size: 12px;
                    }}
                """)
            else:
                btn.setFixedHeight(28)
                btn.setStyleSheet(f"""
                    QPushButton {{
                        background-color: {panel_alt};
                        color: {muted};
                        font-weight: 500;
                        border-top: 1px solid {line};
                        border-left: {border_left};
                        border-right: 1px solid {line};
                        border-bottom: 1px solid {line};
                        border-top-left-radius: {tl_radius};
                        border-top-right-radius: {tr_radius};
                        border-bottom-left-radius: 0px;
                        border-bottom-right-radius: 0px;
                        padding: 4px 16px;
                        font-size: 12px;
                    }}
                    QPushButton:hover {{
                        background-color: {panel_alt_hover};
                        color: {text};
                    }}
                """)

        if hasattr(self, "tabs_widget"):
            self.tabs_widget.raise_()

    def _set_active_tab(self, tab: str) -> None:
        """상태 탭 전환 (전체/등록/진행/완료) 및 스타일 갱신"""
        self._current_tab = tab
        self._update_tab_styles()
        self._apply_filters()

    def _on_period_combo_changed(self) -> None:
        """기간 콤보박스 변경 시 시작일/종료일 자동 설정 및 필터링"""
        selected = self.period_combo.currentText()
        today = date.today()
        self.start_date_edit.blockSignals(True)
        self.end_date_edit.blockSignals(True)

        if selected == "오늘":
            self.start_date_edit.setDate(today)
            self.end_date_edit.setDate(today)
        elif selected == "이번 주":
            start_w = today - timedelta(days=today.weekday())
            end_w = start_w + timedelta(days=6)
            self.start_date_edit.setDate(start_w)
            self.end_date_edit.setDate(end_w)
        elif selected == "이번 달":
            start_m = date(today.year, today.month, 1)
            if today.month == 12:
                end_m = date(today.year, 12, 31)
            else:
                end_m = date(today.year, today.month + 1, 1) - timedelta(days=1)
            self.start_date_edit.setDate(start_m)
            self.end_date_edit.setDate(end_m)
        elif selected == "최근 3개월":
            self.start_date_edit.setDate(today - timedelta(days=90))
            self.end_date_edit.setDate(today)
        elif selected == "올해":
            self.start_date_edit.setDate(date(today.year, 1, 1))
            self.end_date_edit.setDate(date(today.year, 12, 31))

        # 전체 기간이더라도 날짜 수정은 항상 가능
        self.start_date_edit.setEnabled(True)
        self.end_date_edit.setEnabled(True)

        self.start_date_edit.blockSignals(False)
        self.end_date_edit.blockSignals(False)
        self._apply_filters()

    def _on_custom_date_changed(self) -> None:
        """사용자가 날짜를 직접 변경 시 '직접 지정' 모드로 전환 후 필터 적용"""
        if self.period_combo.currentText() != "직접 지정":
            self.period_combo.blockSignals(True)
            idx = self.period_combo.findText("직접 지정")
            if idx >= 0:
                self.period_combo.setCurrentIndex(idx)
            self.period_combo.blockSignals(False)
        self._apply_filters()

    def _open_category_manager(self) -> None:
        cats = get_task_categories(self.repository)
        dlg = SimpleListManagerDialog(
            self,
            "업무 분류 항목 관리",
            cats,
            lambda new_cats: save_task_categories(self.repository, new_cats),
            immutable_item="일반",
            palette=self.palette,
        )
        if dlg.exec() == QDialog.Accepted:
            self._refresh_filter_categories()
            self._apply_filters()

    def _open_status_manager(self) -> None:
        statuses = get_task_statuses(self.repository)
        dlg = SimpleListManagerDialog(
            self,
            "처리 상태 항목 관리",
            statuses,
            lambda new_statuses: save_task_statuses(self.repository, new_statuses),
            immutable_item="등록",
            palette=self.palette,
        )
        if dlg.exec() == QDialog.Accepted:
            self._refresh_filter_statuses()
            self._apply_filters()

    def reload_tasks(self) -> None:
        """저장소에서 모든 TASK 타입 엔트리를 불러와 새로고침"""
        all_entries = self.repository.list_all_entries()
        self._all_tasks = [e for e in all_entries if e.entry_type == EntryType.TASK]
        self._all_tasks.sort(key=lambda x: (x.day or date.min, x.created_at or datetime.min), reverse=True)
        self._apply_filters()

    def _apply_filters(self) -> None:
        """탭(전체/등록/진행/완료), 분류, 상태, 기간, 검색어 필터 적용"""
        query = self.search_input.text().strip().lower()
        selected_cat = self.category_combo.currentText()
        selected_status = self.status_combo.currentText()
        selected_period = self.period_combo.currentText()

        filtered = []
        for task in self._all_tasks:
            # 1. 상태 탭 필터 (전체 / 등록 / 진행 / 완료)
            if self._current_tab == "reg" and (task.status or "등록") != "등록":
                continue
            elif self._current_tab == "prog" and task.status not in ("진행", "진행중"):
                continue
            elif self._current_tab == "done" and task.status != "완료":
                continue

            # 2. 분류 필터
            if selected_cat != "분류":
                task_cat = task.memo_group or "일반"
                if task_cat != selected_cat:
                    continue

            # 3. 상태 필터
            if selected_status != "상태":
                if (task.status or "등록") != selected_status:
                    continue

            # 4. 기간 필터 (시작일 ~ 종료일)
            if selected_period != "기간":
                t_day = task.day or (task.created_at.date() if task.created_at else None)
                if not t_day:
                    continue
                start_d = self.start_date_edit.date().toPython()
                end_d = self.end_date_edit.date().toPython()
                if not (start_d <= t_day <= end_d):
                    continue

            # 5. 검색어 필터 (제목, 기안자, 세부내용, 분류)
            if query:
                title_match = query in (task.title or "").lower()
                author_match = query in (task.assignee or "").lower()
                desc_match = query in (task.description or "").lower()
                cat_match = query in (task.memo_group or "").lower()
                if not (title_match or author_match or desc_match or cat_match):
                    continue

            filtered.append(task)

        self._filtered_tasks = filtered
        self._sort_tasks()
        self._render_table()

    def _update_header_labels(self) -> None:
        """현재 정렬 컬럼 및 방향에 따라 헤더 라벨에 정렬 화살표(▲/▼) 표시"""
        labels = []
        for idx, name in enumerate(self.HEADER_COLUMNS):
            if idx == self._sort_column:
                arrow = " ▲" if self._sort_order == Qt.AscendingOrder else " ▼"
                labels.append(f"{name}{arrow}")
            else:
                labels.append(name)
        self.table.setHorizontalHeaderLabels(labels)

    def _on_header_clicked(self, col: int) -> None:
        """컬럼 헤더 클릭 시 해당 열 기준 오름차순/내림차순 정렬 (편집 컬럼 제외)"""
        if col >= 6:
            return

        if self._sort_column == col:
            self._sort_order = Qt.AscendingOrder if self._sort_order == Qt.DescendingOrder else Qt.DescendingOrder
        else:
            self._sort_column = col
            # 등록일자는 최신순(내림차순) 기본, 기타 항목은 오름차순 기본
            self._sort_order = Qt.DescendingOrder if col == 0 else Qt.AscendingOrder

        self._update_header_labels()
        self._sort_tasks()
        self._render_table()

    def _on_section_resized(self, logicalIndex: int, oldSize: int, newSize: int) -> None:
        """사용자가 헤더를 좌우로 드래그하여 컬럼 크기를 조절할 때 호출 (디바운스 저장)"""
        if getattr(self, "_restoring_columns", False):
            return
        if hasattr(self, "_save_columns_timer"):
            self._save_columns_timer.start()

    def _save_column_widths(self) -> None:
        """조정된 컬럼 너비를 저장소에 영구 저장"""
        try:
            widths = [self.table.columnWidth(i) for i in range(self.table.columnCount())]
            self.repository.set_setting("task_table_column_widths", json.dumps(widths))
            self.repository.save()
        except Exception:
            pass

    def _restore_column_widths(self) -> None:
        """저장소에 저장된 사용자 정의 컬럼 너비 복원"""
        self._restoring_columns = True
        try:
            raw = self.repository.get_setting("task_table_column_widths")
            widths = None
            if raw:
                widths = json.loads(raw)
            if widths and isinstance(widths, list) and len(widths) == self.table.columnCount():
                for i, w in enumerate(widths):
                    if i == 6:
                        self.table.setColumnWidth(i, 55)
                    else:
                        self.table.setColumnWidth(i, max(40, int(w)))
            else:
                for i, w in enumerate(self._default_col_widths):
                    self.table.setColumnWidth(i, w)
        except Exception:
            for i, w in enumerate(self._default_col_widths):
                self.table.setColumnWidth(i, w)
        finally:
            self._restoring_columns = False

    def closeEvent(self, event) -> None:
        self._save_column_widths()
        super().closeEvent(event)

    def _sort_tasks(self) -> None:
        """현재 self._sort_column 및 self._sort_order에 따라 self._filtered_tasks 정렬"""
        if self._sort_column is None or not self._filtered_tasks:
            return

        reverse = (self._sort_order == Qt.DescendingOrder)

        if self._sort_column == 0:  # 등록일자
            def key_fn(t):
                d = t.day or t.start_date or (t.created_at.date() if t.created_at else date.min)
                dt = t.created_at or datetime.min
                return (d, dt)
        elif self._sort_column == 1:  # 분류
            def key_fn(t):
                return (t.memo_group or "일반", t.day or date.min)
        elif self._sort_column == 2:  # 업무 제목
            def key_fn(t):
                return ((t.title or "").strip().lower(), t.day or date.min)
        elif self._sort_column == 3:  # 기안자 / 작성자
            def key_fn(t):
                return ((t.assignee or "").strip().lower(), t.day or date.min)
        elif self._sort_column == 4:  # 상태
            statuses = get_task_statuses(self.repository)
            status_order = {s: i for i, s in enumerate(statuses)}
            def key_fn(t):
                st = t.status or "등록"
                return (status_order.get(st, 999), st, t.day or date.min)
        elif self._sort_column == 5:  # 비고 / 세부내용
            def key_fn(t):
                return ((t.description or "").strip().lower(), t.day or date.min)
        else:
            return

        self._filtered_tasks.sort(key=key_fn, reverse=reverse)

    def _open_meta_manager_menu(self) -> None:
        """분류/상태 설정 드롭다운 메뉴"""
        menu = QMenu(self)
        menu.setStyleSheet(self._menu_stylesheet())
        act_cat = menu.addAction("📁 업무 분류 항목 관리...")
        act_stat = menu.addAction("🏷️ 처리 상태 항목 관리...")
        chosen = menu.exec(self.btn_manage_meta.mapToGlobal(QPoint(0, self.btn_manage_meta.height() + 2)))
        if chosen == act_cat:
            self._open_category_manager()
        elif chosen == act_stat:
            self._open_status_manager()

    def _show_category_menu(self, task: CalendarEntry, btn: QPushButton) -> None:
        """분류 뱃지 클릭 시 드롭다운 메뉴"""
        cats = get_task_categories(self.repository)
        menu = QMenu(self)
        menu.setStyleSheet(self._menu_stylesheet())
        actions = {}
        for c in cats:
            act = menu.addAction(f"📁 {c}")
            actions[act] = c

        menu.addSeparator()
        act_manage = menu.addAction("⚙️ 분류 항목 관리...")

        chosen = menu.exec(btn.mapToGlobal(QPoint(0, btn.height() + 2)))
        if chosen in actions:
            self._on_category_cell_changed(task, actions[chosen])
            self._render_table()
        elif chosen == act_manage:
            self._open_category_manager()

    def _show_status_menu(self, task: CalendarEntry, btn: QPushButton, row_idx: int) -> None:
        """상태 뱃지 클릭 시 드롭다운 메뉴"""
        statuses = get_task_statuses(self.repository)
        menu = QMenu(self)
        menu.setStyleSheet(self._menu_stylesheet())
        actions = {}
        is_dark = self.palette.get("bg", "").lower() in ("#0a0c10", "#171b22") or self.palette.get("text", "").lower() == "#f3f6fb"
        badge_status_cfg = STATUS_CONFIG_DARK if is_dark else STATUS_CONFIG
        for s in statuses:
            icon = badge_status_cfg.get(s, {}).get("icon", "▫️")
            act = menu.addAction(f"{icon} {s}")
            actions[act] = s

        menu.addSeparator()
        act_manage = menu.addAction("⚙️ 상태 항목 관리...")

        chosen = menu.exec(btn.mapToGlobal(QPoint(0, btn.height() + 2)))
        if chosen in actions:
            self._on_status_cell_changed(task, actions[chosen], row_idx)
            self._render_table()
        elif chosen == act_manage:
            self._open_status_manager()

    def _on_paging_combo_changed(self) -> None:
        """페이징 단위 콤보박스 변경 시 초기화 및 렌더링"""
        self._current_page = 1
        self._current_date_anchor = None
        mode = self.paging_combo.currentText()
        try:
            self.repository.set_setting("task_paging_mode", mode)
            self.repository.save()
        except Exception:
            pass
        self._render_table()

    def _get_paged_tasks(self) -> list[CalendarEntry]:
        """현재 선택된 페이징 모드(전체/갯수별/일별/주간별/월별)에 따라 표시할 항목 슬라이싱"""
        mode = self.paging_combo.currentText() if hasattr(self, "paging_combo") else "전체"
        total_items = len(self._filtered_tasks)

        if mode == "전체" or not mode:
            if hasattr(self, "paging_nav_widget"):
                self.paging_nav_widget.setVisible(False)
            return self._filtered_tasks

        # 1. 갯수별 페이징 (10개씩, 20개씩, 50개씩)
        if mode in ("10개씩", "20개씩", "50개씩"):
            self.paging_nav_widget.setVisible(True)
            page_size = int(mode.replace("개씩", ""))
            total_pages = max(1, (total_items + page_size - 1) // page_size)
            if self._current_page > total_pages:
                self._current_page = total_pages
            if self._current_page < 1:
                self._current_page = 1

            start_idx = (self._current_page - 1) * page_size
            end_idx = min(start_idx + page_size, total_items)
            paged = self._filtered_tasks[start_idx:end_idx]

            self.lbl_page_info.setText(f"{self._current_page} / {total_pages} (총 {total_items}건)")
            self.btn_page_prev.setEnabled(self._current_page > 1)
            self.btn_page_next.setEnabled(self._current_page < total_pages)
            return paged

        # 날짜 기반 페이징 (일별 / 주간별 / 월별)
        distinct_dates = sorted(list({
            t.day or t.start_date or (t.created_at.date() if t.created_at else None)
            for t in self._filtered_tasks
            if (t.day or t.start_date or (t.created_at.date() if t.created_at else None))
        }))

        if getattr(self, "_current_date_anchor", None) is None:
            today = date.today()
            if today in distinct_dates or not distinct_dates:
                self._current_date_anchor = today
            else:
                self._current_date_anchor = distinct_dates[-1]

        anchor = self._current_date_anchor

        # 2. 일별 페이징
        if mode == "일별":
            self.paging_nav_widget.setVisible(True)
            paged = [
                t for t in self._filtered_tasks
                if (t.day or t.start_date or (t.created_at.date() if t.created_at else None)) == anchor
            ]
            self.lbl_page_info.setText(f"{anchor.strftime('%Y-%m-%d')} ({len(paged)}건)")
            prev_dates = [d for d in distinct_dates if d < anchor]
            next_dates = [d for d in distinct_dates if d > anchor]
            self.btn_page_prev.setEnabled(len(prev_dates) > 0)
            self.btn_page_next.setEnabled(len(next_dates) > 0)
            return paged

        # 3. 주간별 페이징 (월~일)
        if mode == "주간별":
            self.paging_nav_widget.setVisible(True)
            start_week = anchor - timedelta(days=anchor.weekday())
            end_week = start_week + timedelta(days=6)
            paged = [
                t for t in self._filtered_tasks
                if (t.day or t.start_date or (t.created_at.date() if t.created_at else None))
                and start_week <= (t.day or t.start_date or (t.created_at.date() if t.created_at else None)) <= end_week
            ]
            self.lbl_page_info.setText(f"{start_week.strftime('%m.%d')}~{end_week.strftime('%m.%d')} ({len(paged)}건)")
            prev_weeks = [d for d in distinct_dates if d < start_week]
            next_weeks = [d for d in distinct_dates if d > end_week]
            self.btn_page_prev.setEnabled(len(prev_weeks) > 0)
            self.btn_page_next.setEnabled(len(next_weeks) > 0)
            return paged

        # 4. 월별 페이징
        if mode == "월별":
            self.paging_nav_widget.setVisible(True)
            year, month = anchor.year, anchor.month
            paged = [
                t for t in self._filtered_tasks
                if (t.day or t.start_date or (t.created_at.date() if t.created_at else None))
                and (t.day or t.start_date or (t.created_at.date() if t.created_at else None)).year == year
                and (t.day or t.start_date or (t.created_at.date() if t.created_at else None)).month == month
            ]
            self.lbl_page_info.setText(f"{year}년 {month:02d}월 ({len(paged)}건)")
            prev_months = [d for d in distinct_dates if (d.year, d.month) < (year, month)]
            next_months = [d for d in distinct_dates if (d.year, d.month) > (year, month)]
            self.btn_page_prev.setEnabled(len(prev_months) > 0)
            self.btn_page_next.setEnabled(len(next_months) > 0)
            return paged

        self.paging_nav_widget.setVisible(False)
        return self._filtered_tasks

    def _on_page_prev(self) -> None:
        """이전 페이지 / 이전 날짜 이동"""
        mode = self.paging_combo.currentText()
        if mode in ("10개씩", "20개씩", "50개씩"):
            if self._current_page > 1:
                self._current_page -= 1
                self._render_table()
        else:
            distinct_dates = sorted(list({
                t.day or t.start_date or (t.created_at.date() if t.created_at else None)
                for t in self._filtered_tasks
                if (t.day or t.start_date or (t.created_at.date() if t.created_at else None))
            }))
            anchor = getattr(self, "_current_date_anchor", date.today())
            if mode == "일별":
                prev_dates = [d for d in distinct_dates if d < anchor]
                if prev_dates:
                    self._current_date_anchor = prev_dates[-1]
                else:
                    self._current_date_anchor -= timedelta(days=1)
                self._render_table()
            elif mode == "주간별":
                start_week = anchor - timedelta(days=anchor.weekday())
                prev_weeks = [d for d in distinct_dates if d < start_week]
                if prev_weeks:
                    self._current_date_anchor = prev_weeks[-1]
                else:
                    self._current_date_anchor -= timedelta(days=7)
                self._render_table()
            elif mode == "월별":
                year, month = anchor.year, anchor.month
                prev_months = [d for d in distinct_dates if (d.year, d.month) < (year, month)]
                if prev_months:
                    self._current_date_anchor = prev_months[-1]
                else:
                    m = month - 1
                    y = year
                    if m < 1:
                        m = 12
                        y -= 1
                    self._current_date_anchor = date(y, m, 1)
                self._render_table()

    def _on_page_next(self) -> None:
        """다음 페이지 / 다음 날짜 이동"""
        mode = self.paging_combo.currentText()
        if mode in ("10개씩", "20개씩", "50개씩"):
            total_items = len(self._filtered_tasks)
            page_size = int(mode.replace("개씩", ""))
            total_pages = max(1, (total_items + page_size - 1) // page_size)
            if self._current_page < total_pages:
                self._current_page += 1
                self._render_table()
        else:
            distinct_dates = sorted(list({
                t.day or t.start_date or (t.created_at.date() if t.created_at else None)
                for t in self._filtered_tasks
                if (t.day or t.start_date or (t.created_at.date() if t.created_at else None))
            }))
            anchor = getattr(self, "_current_date_anchor", date.today())
            if mode == "일별":
                next_dates = [d for d in distinct_dates if d > anchor]
                if next_dates:
                    self._current_date_anchor = next_dates[0]
                else:
                    self._current_date_anchor += timedelta(days=1)
                self._render_table()
            elif mode == "주간별":
                start_week = anchor - timedelta(days=anchor.weekday())
                end_week = start_week + timedelta(days=6)
                next_weeks = [d for d in distinct_dates if d > end_week]
                if next_weeks:
                    self._current_date_anchor = next_weeks[0]
                else:
                    self._current_date_anchor += timedelta(days=7)
                self._render_table()
            elif mode == "월별":
                year, month = anchor.year, anchor.month
                next_months = [d for d in distinct_dates if (d.year, d.month) > (year, month)]
                if next_months:
                    self._current_date_anchor = next_months[0]
                else:
                    m = month + 1
                    y = year
                    if m > 12:
                        m = 1
                        y += 1
                    self._current_date_anchor = date(y, m, 1)
                self._render_table()

    def _render_table(self) -> None:
        """필터링 및 페이징된 작업 목록을 7개 컬럼으로 렌더링하고 인라인 수정 지원"""
        if hasattr(self, "hover_delegate"):
            self.hover_delegate.hovered_row = -1
        self.table.blockSignals(True)
        self.table.setRowCount(0)

        tasks_to_show = self._get_paged_tasks()
        self._displayed_tasks = tasks_to_show
        self.table.setRowCount(len(tasks_to_show))

        self._update_stats_label()

        is_dark = self.palette.get("bg", "").lower() in ("#0a0c10", "#171b22") or self.palette.get("text", "").lower() == "#f3f6fb"
        badge_status_cfg = STATUS_CONFIG_DARK if is_dark else STATUS_CONFIG
        badge_cat_cfg = CATEGORY_CONFIG_DARK if is_dark else CATEGORY_CONFIG
        text_normal = QColor(self.palette.get("text", "#2D3748"))
        text_muted = QColor(self.palette.get("muted", "#A0AEC0"))

        row_font = QFont("Malgun Gothic", 10)
        for row_idx, task in enumerate(tasks_to_show):
            is_done = (task.status == "완료")
            text_color = text_muted if is_done else text_normal

            if is_done:
                f = QFont(row_font)
                f.setStrikeOut(True)
            else:
                f = row_font

            # 0: 등록일자 (직접 편집 가능)
            d_val = task.day or task.start_date or (task.created_at.date() if task.created_at else None)
            d_str = d_val.strftime("%Y-%m-%d") if d_val else "-"
            item_date = QTableWidgetItem(d_str)
            item_date.setTextAlignment(Qt.AlignCenter)
            item_date.setForeground(text_color)
            item_date.setFont(f)
            item_date.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled | Qt.ItemIsEditable)
            self.table.setItem(row_idx, 0, item_date)

            # 1: 분류 (알약 뱃지 버튼 클릭 드롭다운) - 1줄 전체 호버/선택 배경 유지를 위해 QTableWidgetItem 삽입
            item_cat = QTableWidgetItem()
            item_cat.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
            self.table.setItem(row_idx, 1, item_cat)
            cat_str = task.memo_group or "일반"
            cat_widget = PillBadgeWidget(
                cat_str,
                badge_cat_cfg,
                lambda btn, t=task: self._show_category_menu(t, btn),
                on_context=lambda pos, t=task: self._show_context_menu_for_task(t, pos),
            )
            cat_widget.setProperty("row_idx", row_idx)
            cat_widget.btn.setProperty("row_idx", row_idx)
            cat_widget.installEventFilter(self)
            cat_widget.btn.installEventFilter(self)
            self.table.setCellWidget(row_idx, 1, cat_widget)

            # 2: 업무 제목 (직접 편집 가능)
            item_title = QTableWidgetItem(task.title or "")
            item_title.setForeground(text_color)
            item_title.setFont(f)
            item_title.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled | Qt.ItemIsEditable)
            self.table.setItem(row_idx, 2, item_title)

            # 3: 기안자 / 작성자 (직접 편집 가능)
            author_str = task.assignee or ""
            item_author = QTableWidgetItem(author_str)
            item_author.setTextAlignment(Qt.AlignCenter)
            item_author.setForeground(text_color)
            item_author.setFont(f)
            item_author.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled | Qt.ItemIsEditable)
            self.table.setItem(row_idx, 3, item_author)

            # 4: 상태 (알약 뱃지 버튼 클릭 드롭다운) - 1줄 전체 호버/선택 배경 유지를 위해 QTableWidgetItem 삽입
            item_status = QTableWidgetItem()
            item_status.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
            self.table.setItem(row_idx, 4, item_status)
            status_str = task.status or "등록"
            status_widget = PillBadgeWidget(
                status_str,
                badge_status_cfg,
                lambda btn, t=task, r=row_idx: self._show_status_menu(t, btn, r),
                on_context=lambda pos, t=task: self._show_context_menu_for_task(t, pos),
            )
            status_widget.setProperty("row_idx", row_idx)
            status_widget.btn.setProperty("row_idx", row_idx)
            status_widget.installEventFilter(self)
            status_widget.btn.installEventFilter(self)
            self.table.setCellWidget(row_idx, 4, status_widget)

            # 5: 비고 / 세부내용 (직접 편집 가능)
            desc_preview = (task.description or "").replace("\n", "  ").strip()
            item_desc = QTableWidgetItem(desc_preview)
            item_desc.setToolTip(task.description or "")
            item_desc.setForeground(text_color)
            item_desc.setFont(f)
            item_desc.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled | Qt.ItemIsEditable)
            self.table.setItem(row_idx, 5, item_desc)

            # 6: 편집 (아이콘 클릭 시 수정 창 열기) - 1줄 전체 호버/선택 배경 유지를 위해 QTableWidgetItem 삽입
            item_edit = QTableWidgetItem()
            item_edit.setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
            self.table.setItem(row_idx, 6, item_edit)

            btn_edit = QPushButton("✏️")
            btn_edit.setToolTip("업무 상세 수정")
            btn_edit.setCursor(Qt.PointingHandCursor)
            btn_edit.setFixedSize(26, 26)
            panel_alt = self.palette.get("panel_alt", "#F8FAFC")
            line = self.palette.get("line", "#CBD5E0")
            line_soft = self.palette.get("line_soft", "#E2E8F0")
            btn_edit.setStyleSheet(f"""
                QPushButton {{
                    background: transparent;
                    border: 1px solid transparent;
                    border-radius: 4px;
                    font-size: 13px;
                    padding: 0px;
                }}
                QPushButton:hover {{
                    background-color: {panel_alt};
                    border: 1px solid {line};
                }}
                QPushButton:pressed {{
                    background-color: {line_soft};
                }}
            """)
            btn_edit.clicked.connect(lambda _, t=task: self._edit_task(t))

            edit_container = QWidget()
            edit_container.setAttribute(Qt.WA_TranslucentBackground, True)
            edit_container.setStyleSheet("background: transparent;")
            edit_layout = QHBoxLayout(edit_container)
            edit_layout.setContentsMargins(0, 0, 0, 0)
            edit_layout.setAlignment(Qt.AlignCenter)
            edit_layout.addWidget(btn_edit)

            edit_container.setProperty("row_idx", row_idx)
            btn_edit.setProperty("row_idx", row_idx)
            edit_container.installEventFilter(self)
            btn_edit.installEventFilter(self)

            self.table.setCellWidget(row_idx, 6, edit_container)

            self.table.setRowHeight(row_idx, 38)

        self.table.blockSignals(False)

    def _update_stats_label(self) -> None:
        """상태 탭 카운트 갱신 (숫자 있는 것만 표기)"""
        total_cnt = len(self._all_tasks)
        reg_cnt = sum(1 for t in self._all_tasks if (t.status or "등록") == "등록")
        prog_cnt = sum(1 for t in self._all_tasks if t.status in ("진행", "진행중"))
        done_cnt = sum(1 for t in self._all_tasks if t.status == "완료")

        self.btn_tab_all.setText(f"전체  {total_cnt}" if total_cnt > 0 else "전체")
        self.btn_tab_reg.setText(f"등록  {reg_cnt}")
        self.btn_tab_prog.setText(f"진행  {prog_cnt}")
        self.btn_tab_done.setText(f"완료  {done_cnt}")

        # 숫자(건수)가 있는 탭만 표기 (전체는 기본 표시)
        self.btn_tab_all.setVisible(True)
        self.btn_tab_reg.setVisible(reg_cnt > 0)
        self.btn_tab_prog.setVisible(prog_cnt > 0)
        self.btn_tab_done.setVisible(done_cnt > 0)

        # 현재 활성화된 탭이 0건이 되어 숨겨진 경우 전체 탭으로 자동 전환
        if self._current_tab == "reg" and reg_cnt == 0:
            self._set_active_tab("all")
        elif self._current_tab == "prog" and prog_cnt == 0:
            self._set_active_tab("all")
        elif self._current_tab == "done" and done_cnt == 0:
            self._set_active_tab("all")
        else:
            self._update_tab_styles()

    def _update_row_appearance(self, row: int, task: CalendarEntry) -> None:
        """완료 상태 변경 시 해당 행 글자 스타일(취소선/색상) 즉시 갱신"""
        is_done = (task.status == "완료")
        text_color = QColor(self.palette.get("muted", "#A0AEC0")) if is_done else QColor(self.palette.get("text", "#2D3748"))
        row_font = QFont("Malgun Gothic", 10)
        f = QFont(row_font)
        if is_done:
            f.setStrikeOut(True)
        for col in (0, 2, 3, 5):
            it = self.table.item(row, col)
            if it:
                it.setFont(f)
                it.setForeground(text_color)

    def _parse_user_date(self, text: str) -> date | None:
        """사용자가 입력한 문자열을 날짜 객체로 변환 (YYYY-MM-DD, YYYY.MM.DD, YYYYMMDD 등 지원)"""
        clean = text.strip().replace(".", "-").replace("/", "-")
        for fmt in ("%Y-%m-%d", "%Y%m%d", "%y-%m-%d"):
            try:
                return datetime.strptime(clean, fmt).date()
            except ValueError:
                pass
        return None

    def _on_table_item_changed(self, item: QTableWidgetItem) -> None:
        """테이블 셀 내에서 직접 텍스트 수정 시 실시간 저장"""
        row = item.row()
        col = item.column()
        if row < 0 or row >= len(self._displayed_tasks):
            return
        if col not in (0, 2, 3, 5):
            return

        task = self._displayed_tasks[row]
        new_val = item.text().strip()
        now = datetime.now()

        if col == 0:  # 등록일자
            parsed = self._parse_user_date(new_val)
            if parsed is None:
                self.table.blockSignals(True)
                d_val = task.day or task.start_date or (task.created_at.date() if task.created_at else None)
                item.setText(d_val.strftime("%Y-%m-%d") if d_val else "-")
                self.table.blockSignals(False)
                QMessageBox.warning(self, "날짜 형식 오류", "날짜는 'YYYY-MM-DD' 형식으로 입력해주세요.\n(예: 2026-09-13)")
                return
            task.day = parsed
            task.start_date = parsed
            if task.created_at:
                task.created_at = datetime.combine(parsed, task.created_at.time())
            else:
                task.created_at = datetime.combine(parsed, now.time())
            task.updated_at = now
            self.table.blockSignals(True)
            item.setText(parsed.strftime("%Y-%m-%d"))
            self.table.blockSignals(False)

        elif col == 2:  # 업무 제목
            if not new_val:
                self.table.blockSignals(True)
                item.setText(task.title or "")
                self.table.blockSignals(False)
                QMessageBox.warning(self, "입력 오류", "업무 제목은 비워둘 수 없습니다.")
                return
            task.title = new_val
            task.updated_at = now

        elif col == 3:  # 기안자 / 작성자
            task.assignee = new_val
            task.updated_at = now

        elif col == 5:  # 비고 / 세부내용
            task.description = item.text()
            item.setToolTip(task.description)
            task.updated_at = now

        self.repository.upsert_entry(task)
        self.repository.save()
        if self.main_window and hasattr(self.main_window, "refresh"):
            self.main_window.refresh()

    def _on_category_cell_changed(self, task: CalendarEntry, new_cat: str) -> None:
        """셀 내 드롭다운에서 분류 변경 시 즉시 저장"""
        if task.memo_group == new_cat:
            return
        task.memo_group = new_cat
        task.updated_at = datetime.now()
        self.repository.upsert_entry(task)
        self.repository.save()
        if self.main_window and hasattr(self.main_window, "refresh"):
            self.main_window.refresh()

    def _on_status_cell_changed(self, task: CalendarEntry, new_status: str, row_idx: int) -> None:
        """셀 내 드롭다운에서 상태 변경 시 즉시 저장 및 스타일 반영"""
        if task.status == new_status:
            return
        task.status = new_status
        task.updated_at = datetime.now()
        self.repository.upsert_entry(task)
        self.repository.save()
        if self.main_window and hasattr(self.main_window, "refresh"):
            self.main_window.refresh()
        self._update_stats_label()
        self._update_row_appearance(row_idx, task)

    def _show_context_menu(self, pos: QPoint) -> None:
        """테이블 우클릭 시 컨텍스트 메뉴 표시"""
        row = self.table.rowAt(pos.y())
        if row < 0 or row >= len(self._displayed_tasks):
            return
        task = self._displayed_tasks[row]
        global_pos = self.table.viewport().mapToGlobal(pos)
        self._show_context_menu_for_task(task, global_pos)

    def _show_context_menu_for_task(self, task: CalendarEntry, global_pos: QPoint) -> None:
        """단일 업무에 대한 컨텍스트 메뉴(상세수정, 출처열기, 삭제 등) 팝업"""
        menu = QMenu(self)
        menu.setStyleSheet(self._menu_stylesheet())

        act_edit = menu.addAction("✏️ 상세 수정 창 열기")
        menu.addSeparator()

        urls = re.findall(r'https?://[^\s<>"]+|www\.[^\s<>"]+', task.description or "")
        act_url = None
        if urls:
            act_url = menu.addAction("🔗 출처 웹페이지 열기")

        act_del = menu.addAction("🗑️ 업무 삭제")

        action = menu.exec(global_pos)
        if action == act_edit:
            self._edit_task(task)
        elif act_url and action == act_url:
            url_str = urls[0]
            if not url_str.startswith("http"):
                url_str = "https://" + url_str
            QDesktopServices.openUrl(QUrl(url_str))
        elif action == act_del:
            self._delete_task(task)

    def _on_add_task(self) -> None:
        """새 업무 등록 (전용 TaskEditDialog 호출)"""
        self._edit_task(None)

    def _edit_task(self, task: CalendarEntry | None) -> None:
        """업무 상세 편집 다이얼로그 열기 (기안자 자동완성 포함)"""
        known_authors = [t.assignee for t in self._all_tasks if t.assignee]
        dlg = TaskEditDialog(self, self.repository, task, known_authors, palette=self.palette)
        if dlg.exec() == QDialog.Accepted:
            self.repository.save()
            if self.main_window and hasattr(self.main_window, "refresh"):
                self.main_window.refresh()
            self._refresh_filter_categories()
            self._refresh_filter_statuses()
            self.reload_tasks()

    def _delete_task(self, task: CalendarEntry) -> None:
        """단일 업무 삭제 확인 및 수행"""
        reply = QMessageBox.question(
            self,
            "삭제 확인",
            f"'{task.title}' 업무를 삭제하시겠습니까?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No,
        )
        if reply == QMessageBox.Yes and task.entry_id:
            self.repository.delete_entry(task.entry_id)
            self.repository.save()
            if self.main_window and hasattr(self.main_window, "refresh"):
                self.main_window.refresh()
            self.reload_tasks()

    def _on_delete_selected(self) -> None:
        """선택된 업무 삭제 버튼 핸들러"""
        selected_row = self.table.currentRow()
        if selected_row < 0 or selected_row >= len(self._filtered_tasks):
            QMessageBox.information(self, "안내", "삭제할 업무를 선택해 주세요.")
            return

        task = self._filtered_tasks[selected_row]
        self._delete_task(task)

    def _export_to_excel(self) -> None:
        """현재 필터링된 업무 목록을 깔끔한 서식의 엑셀 파일로 내보내기"""
        if not self._filtered_tasks:
            QMessageBox.information(self, "안내", "내보낼 업무 항목이 없습니다.")
            return

        now_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"업무관리대장_{now_str}.xlsx"

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "업무 목록 엑셀 내보내기",
            default_name,
            "Excel Files (*.xlsx)",
        )
        if not file_path:
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "업무관리대장"

            headers = ["번호", "상태", "등록일자", "업무분류", "업무제목", "기안자/작성자", "세부내용/출처"]
            ws.append(headers)

            header_fill = PatternFill(start_color="4A5568", end_color="4A5568", fill_type="solid")
            header_font = Font(name="Malgun Gothic", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style="thin", color="E2E8F0"),
                right=Side(style="thin", color="E2E8F0"),
                top=Side(style="thin", color="E2E8F0"),
                bottom=Side(style="thin", color="E2E8F0"),
            )

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
            ws.row_dimensions[1].height = 28

            for idx, task in enumerate(self._filtered_tasks, 1):
                is_done = (task.status == "완료")
                d_val = task.day or task.start_date or (task.created_at.date() if task.created_at else None)
                d_str = d_val.strftime("%Y-%m-%d") if d_val else ""
                cat_str = task.memo_group or "일반"
                author_str = task.assignee or ""
                status_str = task.status or "등록"
                desc_str = task.description or ""

                row_values = [
                    idx,
                    status_str,
                    d_str,
                    cat_str,
                    task.title,
                    author_str,
                    desc_str,
                ]
                ws.append(row_values)
                row_idx = idx + 1

                for col_idx in range(1, len(headers) + 1):
                    c = ws.cell(row=row_idx, column=col_idx)
                    c.font = Font(name="Malgun Gothic", size=10)
                    c.border = thin_border
                    c.alignment = Alignment(vertical="center")

                    if col_idx in (1, 2, 3, 6):
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_idx == 4:
                        c.alignment = Alignment(horizontal="center", vertical="center")
                        c.font = Font(name="Malgun Gothic", size=10, bold=True, color="6C5CE7" if not is_done else "718096")

                    if is_done and col_idx == 5:
                        c.font = Font(name="Malgun Gothic", size=10, strike=True, color="A0AEC0")

                ws.row_dimensions[row_idx].height = 22

            col_widths = {1: 8, 2: 12, 3: 14, 4: 12, 5: 35, 6: 18, 7: 40}
            for col_idx, width in col_widths.items():
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = width

            wb.save(file_path)

            reply = QMessageBox.question(
                self,
                "엑셀 저장 완료",
                f"총 {len(self._filtered_tasks)}건의 업무 목록이 성공적으로 저장되었습니다.\n\n파일을 지금 열어보시겠습니까?",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.Yes,
            )
            if reply == QMessageBox.Yes:
                os.startfile(file_path)

        except Exception as e:
            logger.exception("Task excel export failed")
            QMessageBox.critical(self, "저장 실패", f"엑셀 파일 저장 중 오류가 발생했습니다:\n{e}")
